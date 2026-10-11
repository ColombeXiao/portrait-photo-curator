#!/usr/bin/env python3
"""Codex portrait-photo-curator deterministic file/report pipeline.

This script NEVER grades photos. Visual review is the responsibility of Codex.
Prepare -> manually populate _work/reviews/*.json -> verify -> finalize -> verify.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any

PHOTO_EXT = {
    '.jpg', '.jpeg', '.jfif', '.png', '.webp', '.bmp', '.tif', '.tiff',
    '.heic', '.heif', '.dng', '.cr2', '.cr3', '.nef', '.arw', '.raf',
    '.rw2', '.orf', '.pef', '.srw', '.3fr', '.iiq', '.erf', '.k25'
}
RAW_EXT = {'.dng', '.cr2', '.cr3', '.nef', '.arw', '.raf', '.rw2', '.orf', '.pef', '.srw', '.3fr', '.iiq', '.erf', '.k25'}
HEIF_EXT = {'.heic', '.heif'}
GRADE_NAMES = {'A': 'A_可直接出片', 'B': 'B_后期可出片', 'C': 'C_废片'}
REPORT_NAMES = {
    'A': 'A档_摄影优点与复现方法',
    'B': 'B档_醒图修图步骤与摄影规则',
}
MAX_REPORT_ITEMS = {'A': 30, 'B': 20}


def fail(msg: str):
    raise ValueError(msg)


def load_json(path: Path):
    with path.open('r', encoding='utf-8') as f:
        return json.load(f)


def save_json(path: Path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + '.tmp')
    with tmp.open('w', encoding='utf-8') as f:
        json.dump(obj, f, ensure_ascii=False, indent=2)
        f.write('\n')
    os.replace(tmp, path)


def safe_name(name: str, limit: int = 110) -> str:
    """Trim filename but preserve extension; P0001 makes names unique."""
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f]', '_', name).strip('. ')
    suffix = Path(name).suffix
    stem = name[:-len(suffix)] if suffix else name
    if not stem:
        stem = 'image'
    return stem[:max(1, limit - len(suffix))] + suffix


def is_below(child: Path, parent: Path) -> bool:
    try:
        child.resolve().relative_to(parent.resolve())
        return True
    except ValueError:
        return False


def source_file(root: Path, rel: str) -> Path:
    p = (root / rel).resolve()
    if not is_below(p, root) or not p.is_file():
        fail(f'照片源路径无效、越界或已消失: {rel}')
    return p


def hash_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda: f.read(2 * 1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def pillow_image(path: Path):
    from PIL import Image, ImageCms, ImageOps
    if path.suffix.lower() in HEIF_EXT:
        try:
            import pillow_heif
            pillow_heif.register_heif_opener()
        except ImportError as e:
            raise RuntimeError('缺少 pillow-heif: python3 -m pip install pillow-heif') from e
    if path.suffix.lower() in RAW_EXT:
        try:
            import rawpy
        except ImportError as e:
            raise RuntimeError('缺少 rawpy: python3 -m pip install rawpy') from e
        try:
            with rawpy.imread(str(path)) as raw:
                arr = raw.postprocess(use_camera_wb=True, output_bps=8)
            return Image.fromarray(arr), {}
        except Exception as e:
            raise RuntimeError(f'RAW 解码失败（可能不支持该型号/文件已损坏）: {e}') from e
    with Image.open(path) as im:
        meta = {}
        try:
            exif = im.getexif()
            def numeric(tag):
                val = exif.get(tag)
                try:
                    return float(val)
                except (ValueError, TypeError, ZeroDivisionError):
                    return None
            # Only read verified source metadata, never interpolate a missing number.
            for tag, label in [(34855, 'ISO'), (37386, 'focal_mm'), (33437, 'fnumber')]:
                val = numeric(tag)
                if val is not None:
                    meta[label] = round(val, 3)
            shutter = numeric(33434)
            if shutter is not None and shutter > 0:
                meta['shutter_seconds'] = round(shutter, 6)
            model = exif.get(272)
            if model:
                meta['camera_model'] = str(model)[:100]
        except Exception:
            pass
        icc = im.info.get('icc_profile')
        fixed = ImageOps.exif_transpose(im)
        fixed.load()
        if icc:
            try:
                import io
                srgb = ImageCms.createProfile('sRGB')
                profile = ImageCms.ImageCmsProfile(io.BytesIO(icc))
                fixed = ImageCms.profileToProfile(fixed.convert('RGB'), profile, srgb, outputMode='RGB')
            except Exception:
                fixed = fixed.convert('RGB')
        elif fixed.mode in ('RGBA', 'LA', 'P'):
            fixed = fixed.convert('RGBA')
            background = Image.new('RGB', fixed.size, (255, 255, 255))
            background.paste(fixed, mask=fixed.getchannel('A'))
            fixed = background
        else:
            fixed = fixed.convert('RGB')
        return fixed, meta


def create_preview(src: Path, dest: Path, longest: int):
    from PIL import Image
    im, meta = pillow_image(src)
    if im.width * im.height > 250_000_000:
        raise RuntimeError('图片像素过大（> 2.5 亿像素），为避免内存风险请单独处理')
    im.thumbnail((longest, longest), Image.Resampling.LANCZOS)
    dest.parent.mkdir(parents=True, exist_ok=True)
    im.save(dest, 'JPEG', quality=89, optimize=True, subsampling=0)
    return {'width': im.width, 'height': im.height, 'exif': meta}


def font_for_sheets(sz: int):
    from PIL import ImageFont
    fonts = [
        '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
        '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
        '/System/Library/Fonts/PingFang.ttc',
        '/Library/Fonts/Arial Unicode.ttf',
    ]
    for p in fonts:
        try:
            return ImageFont.truetype(p, size=sz)
        except OSError:
            pass
    return ImageFont.load_default()


def draw_contact_sheet(batch, base: Path, dest: Path):
    from PIL import Image, ImageDraw, ImageOps
    cols, cw, ch, pad = 4, 270, 242, 12
    rows = (len(batch) + cols - 1) // cols
    canvas = Image.new('RGB', (cols * cw, rows * ch), (240, 241, 243))
    d = ImageDraw.Draw(canvas)
    big_font = font_for_sheets(17)
    small_font = font_for_sheets(12)
    for n, photo in enumerate(batch):
        x, y = (n % cols) * cw, (n // cols) * ch
        im = Image.open(base / photo['preview']).convert('RGB')
        thumb = ImageOps.contain(im, (cw - 2 * pad, 185), Image.Resampling.LANCZOS)
        canvas.paste(thumb, (x + (cw-thumb.width)//2, y + 5 + (185-thumb.height)//2))
        d.text((x + 10, y + 193), photo['id'], font=big_font, fill=(35, 35, 39))
        basename = Path(photo['relative_path']).name
        d.text((x + 10, y + 218), basename[:32], font=small_font, fill=(70, 70, 75))
    dest.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(dest, 'JPEG', quality=90, optimize=True)


def prepare(args):
    try:
        import PIL, docx  # noqa - detect dependencies early
    except ImportError:
        fail('缺少依赖。先执行 python3 -m pip install -r <skill目录>/requirements.txt')
    root = Path(args.input).expanduser().resolve()
    output = Path(args.output).expanduser().resolve()
    if not root.is_dir():
        fail(f'输入文件夹不存在: {root}')
    if is_below(output, root) or root == output:
        fail('为防扫描混入结果和误改原件，输出目录必须在输入目录之外')
    if (output / '_work' / 'manifest.json').exists():
        fail(f'此输出目录已准备过: {output}。为安全起见使用新输出目录。')
    if output.exists() and any(output.iterdir()):
        fail(f'输出目录不为空，为避免覆盖已有内容请使用新目录: {output}')
    photos = sorted((p for p in root.rglob('*') if p.is_file() and not p.is_symlink()
                     and p.suffix.lower() in PHOTO_EXT
                     and not any(part.startswith('.') for part in p.relative_to(root).parts)),
                    key=lambda p: str(p.relative_to(root)).casefold())
    if not photos:
        fail('输入目录没有扫描到受支持的图像，检查照片路径和扩展名')
    missing = []
    if any(p.suffix.lower() in HEIF_EXT for p in photos):
        try: import pillow_heif  # noqa: F401
        except ImportError: missing.append('pillow-heif（HEIC/HEIF）')
    if any(p.suffix.lower() in RAW_EXT for p in photos):
        try: import rawpy  # noqa: F401
        except ImportError: missing.append('rawpy（相机 RAW）')
    if missing:
        fail('本组照片需要额外安装依赖：' + '、'.join(missing) + '\n例如 python3 -m pip install pillow-heif rawpy')
    work = output / '_work'
    records = []
    for index, path in enumerate(photos, start=1):
        photo_id = f'P{index:04d}'
        rel = path.relative_to(root).as_posix()
        rec = {'id': photo_id, 'relative_path': rel,
               'preview': f'_work/previews/{photo_id}.jpg',
               'copy_name': f'{photo_id}__{safe_name(path.name)}',
               'size_bytes': path.stat().st_size}
        try:
            rec.update(create_preview(path, output / rec['preview'], args.preview_size))
            rec['preview_status'] = 'ok'
        except Exception as e:
            rec['preview_status'] = 'error'
            rec['preview_error'] = str(e)
            print(f'无法预览 {rel}: {e}', file=sys.stderr)
        records.append(rec)
        if index % 30 == 0:
            print(f'已准备预览 {index}/{len(photos)}')
    manifest = {'schema_version': 1, 'source_root': str(root), 'output_root': str(output),
                'batch_size': args.batch_size, 'photos': records}
    save_json(work / 'manifest.json', manifest)
    if all(rec['preview_status'] == 'ok' for rec in records):
        for i in range(0, len(records), args.batch_size):
            num = i // args.batch_size + 1
            batch = records[i:i+args.batch_size]
            fname = f'batch_{num:03d}'
            # Keep each batch very small and human-readable for an LLM.
            save_json(work / 'batches' / f'{fname}.json', {
                'batch_name': fname,
                'items': [{k: item[k] for k in ('id', 'relative_path', 'preview')}
                          for item in batch]
            })
            draw_contact_sheet(batch, output, work / 'contact_sheets' / f'{fname}.jpg')
        (work / 'reviews').mkdir(parents=True, exist_ok=True)
    print(f'发现 {len(records)} 张照片；预览成功 {sum(r["preview_status"]=="ok" for r in records)} 张。')
    print(f'输出: {output}\n分批清单: {work / "batches"}\n审片工作目录: {work / "reviews"}')
    if any(r['preview_status'] != 'ok' for r in records):
        fail('部分照片无法解码，没有生成可用批次；修复依赖/文件后使用新输出目录重新 prepare，禁止遗漏。')


def required_text(obj, key, context):
    if not isinstance(obj, dict) or not isinstance(obj.get(key), str) or not obj[key].strip():
        fail(f'{context}: 缺少非空文本 {key}')


def list_of_text(obj, key, context):
    rows = obj.get(key)
    if not isinstance(rows, list) or not rows or any(not isinstance(x, str) or not x.strip() for x in rows):
        fail(f'{context}: {key} 必须是至少一条非空中文文本列表')


def list_of_dicts(obj, key, context, fields: tuple, min_count=1):
    rows = obj.get(key)
    if not isinstance(rows, list) or len(rows) < min_count:
        fail(f'{context}: {key} 至少 {min_count} 条')
    for n, row in enumerate(rows, 1):
        for field in fields:
            required_text(row, field, f'{context} {key}[{n}]')


def validate_review(row, known_ids):
    if not isinstance(row, dict):
        fail('评审项必须是对象')
    photo_id = row.get('id')
    if photo_id not in known_ids:
        fail(f'评审项有不存在的 ID: {photo_id}')
    g = row.get('grade')
    context = f'{photo_id} / {g}'
    if g not in GRADE_NAMES:
        fail(f'{context}: grade 必须是 A/B/C')
    if g == 'C':
        if set(row) != {'id', 'grade'}:
            fail(f'{context}: C 档只能含 id、grade，不能有照片级描述')
        return
    required_text(row, 'title', context)
    if g == 'A':
        required_text(row, 'verdict', context)
        list_of_dicts(row, 'strengths', context, ('observation', 'principle', 'repeat'), 3)
        list_of_text(row, 'next_shoot', context)
    else:
        required_text(row, 'potential', context)
        list_of_dicts(row, 'issues', context,
                      ('observation', 'impact', 'principle', 'capture_prevention'))
        xt = row.get('xingtu')
        required_text(xt, 'goal', context)
        required_text(xt, 'export', context)
        list_of_dicts(xt, 'steps', context, ('tool', 'action', 'start', 'reason', 'check'))
        list_of_text(xt, 'avoid', context)
        list_of_dicts(row, 'learning', context, ('rule', 'applied', 'next_shoot'))
        list_of_text(row, 'reshoot', context)


def collect_reviews(output: Path):
    manifest_file = output / '_work' / 'manifest.json'
    if not manifest_file.is_file():
        fail('未找到 _work/manifest.json。请先 prepare')
    manifest = load_json(manifest_file)
    if manifest.get('schema_version') != 1:
        fail('manifest 版本不受支持')
    if Path(manifest['output_root']).resolve() != output.resolve():
        fail('manifest 输出路径与当前 --output 不一致')
    src_root = Path(manifest['source_root']).resolve()
    if not src_root.is_dir():
        fail('原始照片目录不在原位置，请恢复原目录后重试')
    rows = manifest['photos']
    if not rows:
        fail('没有照片')
    ids = {r['id'] for r in rows}
    if len(ids) != len(rows):
        fail('manifest ID 重复')
    bad_previews = [p['id'] for p in rows if p.get('preview_status') != 'ok'
                    or not (output / p['preview']).is_file()]
    if bad_previews:
        fail('无法预览的照片，不能分级：' + ', '.join(bad_previews))
    batches = sorted((output / '_work' / 'batches').glob('batch_*.json'))
    expected_batch_count = (len(rows) + manifest['batch_size'] - 1) // manifest['batch_size']
    if len(batches) != expected_batch_count:
        fail('缺少分批清单，不能继续')
    mapping = {}
    for batch_file in batches:
        batch = load_json(batch_file)
        expect = [i['id'] for i in batch.get('items', [])]
        review_file = output / '_work' / 'reviews' / batch_file.name
        if not review_file.is_file():
            fail(f'尚未完成审片: {review_file.name}（本批 {len(expect)} 张）')
        parsed = load_json(review_file)
        items = parsed.get('items')
        if not isinstance(items, list):
            fail(f'{review_file.name}: 必须是 items 数组')
        found = [item.get('id') if isinstance(item, dict) else None for item in items]
        if len(items) != len(expect) or len(set(found)) != len(found) or set(found) != set(expect):
            fail(f'{review_file.name}: 应有 {expect}，实有 {found}；不可漏掉、多出或重复')
        for item in items:
            validate_review(item, ids)
            if item['id'] in mapping:
                fail(f'跨批次出现重复照片 {item["id"]}')
            mapping[item['id']] = item
    if set(mapping) != ids:
        fail('实际评审与原照片清单数量不一致')
    counts = {g: sum(1 for x in mapping.values() if x['grade'] == g) for g in GRADE_NAMES}
    return manifest, mapping, counts


def set_font(run, size=None, bold=False, color=None):
    from docx.shared import Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run.font.name = 'Noto Sans CJK SC'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    if size:
        run.font.size = Pt(size)
    if bold:
        run.bold = True
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def para(document, label: str, body: str, style=None, keep_next=False):
    p = document.add_paragraph(style=style)
    p.paragraph_format.keep_with_next = keep_next
    if label:
        set_font(p.add_run(label), 10.5, bold=True)
    set_font(p.add_run(body), 10.5)
    return p


def section(document, title: str):
    p = document.add_paragraph()
    p.paragraph_format.space_before = __import__('docx').shared.Pt(9)
    set_font(p.add_run(title), 12, bold=True, color=(42, 65, 86))
    return p


def make_document(title: str, grade: str, items, manifest, output: Path, dest: Path):
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.9)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    styles = doc.styles
    default = styles['Normal']
    default.font.name = 'Noto Sans CJK SC'
    default.font.size = Pt(10.5)
    default.paragraph_format.space_after = Pt(5)
    default.paragraph_format.line_spacing = 1.23
    set_font(doc.add_heading(title, 0).runs[0], 20, True)
    intro = ('A 档为可直接展示作品；下述优点均应对应照片中的可见证据，帮助下一次稳定复现。'
             if grade == 'A' else
             'B 档为可通过后期合理改善的作品。醒图参数仅是起点，各版本滑杆和功能不同；以自然肤色、视觉层级和原片对比为准。')
    para(doc, '', intro)
    para(doc, '审片日期：', datetime.now().strftime('%Y-%m-%d'))
    para(doc, '本册张数：', str(len(items)))
    para(doc, '图片说明：', '文中缩略图仅用于定位，最终判断及复制均以原文件为准；原图未作修改。')
    for j, (photo, review) in enumerate(items):
        if j > 0:
            doc.add_page_break()
        doc.add_heading(f'【{photo["id"]}】 {review["title"]}', level=1)
        para(doc, '源文件：', photo['relative_path'])
        para(doc, '归档照片：', f'{GRADE_NAMES[grade]}/照片/{photo["copy_name"]}')
        exif = photo.get('exif') or {}
        if exif:
            parts = []
            for k, label in [('camera_model','设备'), ('focal_mm','焦距mm'),
                             ('fnumber','光圈f/'), ('shutter_seconds','快门s'), ('ISO','ISO')]:
                if k in exif:
                    parts.append(f'{label} {exif[k]}')
            if parts:
                para(doc, '文件 EXIF：', ' · '.join(parts))
        prev = output / photo['preview']
        if prev.is_file():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            from PIL import Image
            with Image.open(prev) as pic:
                ratio = pic.height / pic.width
            if ratio >= 1.12:
                p.add_run().add_picture(str(prev), height=Cm(7.4))
            else:
                p.add_run().add_picture(str(prev), width=Cm(10.6))
        if grade == 'A':
            section(doc, '专业判断｜为什么这一张能直接出片')
            para(doc, '', review['verdict'])
            section(doc, '可复现的优秀之处')
            for n, s in enumerate(review['strengths'], 1):
                para(doc, f'{n}. 看得见的优点：', s['observation'], keep_next=True)
                para(doc, '   摄影规则：', s['principle'], keep_next=True)
                para(doc, '   下次保持：', s['repeat'])
            section(doc, '下次拍摄要延续的动作')
            for n, v in enumerate(review['next_shoot'], 1):
                para(doc, f'{n}. ', v)
        else:
            section(doc, '拍摄价值｜保留这张的原因')
            para(doc, '', review['potential'])
            section(doc, '具体不足｜为什么影响成片')
            for n, i in enumerate(review['issues'], 1):
                para(doc, f'{n}. 问题：', i['observation'], keep_next=True)
                para(doc, '   视觉影响：', i['impact'], keep_next=True)
                para(doc, '   摄影原理：', i['principle'], keep_next=True)
                para(doc, '   下次预防：', i['capture_prevention'])
            xt = review['xingtu']
            section(doc, '醒图逐步操作｜对着手机可以一步步做')
            para(doc, '本次目标：', xt['goal'])
            for n, s in enumerate(xt['steps'], 1):
                para(doc, f'第 {n} 步｜{s["tool"]}：', s['action'], keep_next=True)
                para(doc, '   调整起点：', s['start'], keep_next=True)
                para(doc, '   为什么做：', s['reason'], keep_next=True)
                para(doc, '   如何验收：', s['check'])
            section(doc, '自然真实与导出')
            for v in xt['avoid']:
                para(doc, '避免：', v)
            para(doc, '导出：', xt['export'])
            section(doc, '摄影规则课堂｜为什么这样改、下次如何拍')
            for n, l in enumerate(review['learning'], 1):
                para(doc, f'{n}. 摄影法则：', l['rule'], keep_next=True)
                para(doc, '   本图应用：', l['applied'], keep_next=True)
                para(doc, '   实拍练习：', l['next_shoot'])
            section(doc, '下一次拍摄的现场行动清单')
            for n, v in enumerate(review['reshoot'], 1):
                para(doc, f'{n}. ', v)
    # footer
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run('旅行人像摄影复盘 · 照片原件保持不变'), 8.5, color=(90, 90, 90))
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)


def build_summary(counts, output: Path):
    total = sum(counts.values())
    lines = ['旅行人像 A/B/C 专业审片结果',
             f'照片总数：{total}',
             f'A 可直接出片：{counts["A"]}',
             f'B 后期可出片：{counts["B"]}',
             f'C 废片：{counts["C"]}',
             '', 'A/B 文件夹内附 Word 摄影复盘与醒图修片指导，C 仅存原图副本。',
             '全部原图只进行了文件级复制，未修改内容、未删除原文件。',
             '逐张审片原始记录见 _work/reviews/（C 仅有 ID/等级）。',
             '使用 verify --phase final 可重新核对文件哈希和报告索引。', '']
    return '\n'.join(lines)


def finalize(args):
    output = Path(args.output).expanduser().resolve()
    manifest, mapping, counts = collect_reviews(output)
    root = Path(manifest['source_root']).resolve()
    # Preflight all paths and collisions before writing results.
    for grade_dir in GRADE_NAMES.values():
        if (output / grade_dir).exists():
            fail(f'检测到已存在的归档文件夹 {grade_dir}，不允许覆盖。请使用新结果目录。')
    if (output / '结果汇总.txt').exists():
        fail('结果汇总.txt 已存在；拒绝覆盖。')
    photos = manifest['photos']
    for photo in photos:
        src = source_file(root, photo['relative_path'])
        if src.stat().st_size != photo['size_bytes']:
            fail(f'原文件在准备后大小发生变化：{photo["relative_path"]}；请重新 prepare')
    by_grade = {g: [(p, mapping[p['id']]) for p in photos if mapping[p['id']]['grade'] == g]
                for g in GRADE_NAMES}
    with tempfile.TemporaryDirectory(prefix='photo_classification_', dir=output / '_work') as tmp:
        stage = Path(tmp)
        for grade, pairs in by_grade.items():
            image_dir = stage / GRADE_NAMES[grade] / '照片'
            image_dir.mkdir(parents=True, exist_ok=True)
            for photo, review in pairs:
                src = source_file(root, photo['relative_path'])
                dest = image_dir / photo['copy_name']
                if dest.exists():
                    fail('目标文件命名冲突: ' + str(dest))
                shutil.copy2(src, dest)
                if hash_file(src) != hash_file(dest):
                    fail(f'复制文件校验不一致：{photo["id"]}，已停止提交')
        for grade in ('A', 'B'):
            pairs = by_grade[grade]
            for start in range(0, len(pairs), MAX_REPORT_ITEMS[grade]):
                page = start // MAX_REPORT_ITEMS[grade] + 1
                dest = stage / GRADE_NAMES[grade] / (REPORT_NAMES[grade] + f'_第{page:02d}册.docx')
                make_document('A 档｜专业摄影师优点复盘' if grade == 'A' else
                              'B 档｜醒图后期与人像摄影教学',
                              grade, pairs[start:start+MAX_REPORT_ITEMS[grade]],
                              manifest, output, dest)
                # Verify docx can be reopened and lists each expected ID.
                from docx import Document
                d = Document(dest)
                body = '\n'.join(p.text for p in d.paragraphs)
                for photo, _ in pairs[start:start+MAX_REPORT_ITEMS[grade]]:
                    if f'【{photo["id"]}】' not in body:
                        fail(f'Word 报告里找不到图片ID: {photo["id"]}')
        # Only after all copies + reports pass checks, move folders into place.
        for grade_dir in GRADE_NAMES.values():
            os.replace(stage / grade_dir, output / grade_dir)
        (output / '结果汇总.txt').write_text(build_summary(counts, output), encoding='utf-8')
    print(f'归档完成，A={counts["A"]}, B={counts["B"]}, C={counts["C"]}，总计 {len(photos)}')
    print(f'结果目录：{output}\n请继续运行 verify --phase final')


def verify_final(output: Path, manifest, mapping, counts):
    from docx import Document
    root = Path(manifest['source_root']).resolve()
    scanned = {}
    for grade, folder in GRADE_NAMES.items():
        target = output / folder
        if not target.is_dir():
            fail(f'缺少档位文件夹: {folder}')
        photo_dir = target / '照片'
        if not photo_dir.is_dir():
            fail(f'缺少档位照片文件夹: {folder}')
        actual = {p.name for p in photo_dir.iterdir() if p.is_file()}
        expected = {photo['copy_name'] for photo in manifest['photos']
                    if mapping[photo['id']]['grade'] == grade}
        if actual != expected:
            fail(f'{folder} 文件数/名称不一致；缺少 {sorted(expected-actual)}，额外 {sorted(actual-expected)}')
        if grade == 'C':
            forbidden = [p for p in target.rglob('*') if p.is_file() and p.parent != photo_dir]
            if forbidden:
                fail('C 档出现多余文件/个片说明：' + ', '.join(str(x) for x in forbidden))
        else:
            reports = sorted(target.glob('*.docx'))
            expected_report_count = (len(expected) + MAX_REPORT_ITEMS[grade] - 1)//MAX_REPORT_ITEMS[grade]
            if len(reports) != expected_report_count:
                fail(f'{folder} Word 报告数量不一致：预期 {expected_report_count}，实际 {len(reports)}')
            ids_in_docs = []
            for report in reports:
                doc = Document(report)
                txt = '\n'.join(p.text for p in doc.paragraphs)
                ids_in_docs.extend(re.findall(r'【(P\d{4,})】', txt))
            needed = [p['id'] for p in manifest['photos'] if mapping[p['id']]['grade'] == grade]
            if sorted(ids_in_docs) != sorted(needed):
                fail(f'{folder} Word 未完整对应照片：expected={needed} actual={ids_in_docs}')
            extraneous = [p for p in target.rglob('*') if p.is_file()
                          and p.parent != photo_dir and p.suffix.lower() != '.docx']
            if extraneous:
                fail(f'{folder} 出现其他格式文件：{extraneous}')
        scanned[grade] = len(actual)
    if scanned != counts:
        fail('分类结果统计不一致')
    for photo in manifest['photos']:
        src = source_file(root, photo['relative_path'])
        g = mapping[photo['id']]['grade']
        copied = output / GRADE_NAMES[g] / '照片' / photo['copy_name']
        if hash_file(src) != hash_file(copied):
            fail(f'{photo["id"]} 文件哈希校验失败，源文件与归档件不同')
    if not (output / '结果汇总.txt').is_file():
        fail('缺少结果汇总.txt')
    print(f'核验通过：总 {len(manifest["photos"])} 张；A={counts["A"]}, B={counts["B"]}, C={counts["C"]}')
    print('每张恰有一个分档、原件哈希一致、A/B 报告逐张包含、C 无文字报告。')


def verify(args):
    output = Path(args.output).expanduser().resolve()
    manifest, mapping, counts = collect_reviews(output)
    print(f'评审内容结构核验通过：A={counts["A"]}, B={counts["B"]}, C={counts["C"]}，没有漏片或重复分档。')
    if args.phase == 'final':
        verify_final(output, manifest, mapping, counts)



def detail(args):
    """Generate a larger JPEG *viewing copy* when near-eye focus is uncertain."""
    output = Path(args.output).expanduser().resolve()
    manifest_file = output / '_work' / 'manifest.json'
    if not manifest_file.exists():
        fail('缺少准备好的 manifest。请先 prepare')
    manifest = load_json(manifest_file)
    if Path(manifest['output_root']).resolve() != output:
        fail('manifest 的输出路径与当前输出目录不一致')
    photo = next((p for p in manifest['photos'] if p['id'] == args.id), None)
    if photo is None:
        fail(f'未知照片 ID：{args.id}')
    root = Path(manifest['source_root']).resolve()
    src = source_file(root, photo['relative_path'])
    path = output / '_work' / 'detail_previews' / f'{photo["id"]}_detail.jpg'
    result = create_preview(src, path, args.size)
    print(f'高分辨率预览：{path} ({result["width"]}×{result["height"]})')


def main():
    parser = argparse.ArgumentParser(description='旅行人像归档与 Word 报告工具（本程序不做照片分级）')
    sub = parser.add_subparsers(dest='command', required=True)
    prep = sub.add_parser('prepare', help='建立预览、联系表与批次清单，供 Codex 真正看图')
    prep.add_argument('--input', required=True, help='原始照片文件夹（保持只读）')
    prep.add_argument('--output', required=True, help='输出文件夹（必须在输入文件夹外）')
    prep.add_argument('--batch-size', type=int, default=16, help='每批评审张数（推荐 12~20）')
    prep.add_argument('--preview-size', type=int, default=2000, help='预览长边像素，推荐 2000')
    prep.set_defaults(func=prepare)
    fin = sub.add_parser('finalize', help='通过完整审片 JSON 安全复制原件并产生 Word 报告')
    fin.add_argument('--output', required=True)
    fin.set_defaults(func=finalize)
    check = sub.add_parser('verify', help='验证审片数据/最终文件完整性')
    check.add_argument('--output', required=True)
    check.add_argument('--phase', choices=('review', 'final'), default='review')
    check.set_defaults(func=verify)
    high = sub.add_parser('detail', help='为不确定失焦/闭眼等情况生成更大视觉预览')
    high.add_argument('--output', required=True)
    high.add_argument('--id', required=True, help='照片 ID，如 P0003')
    high.add_argument('--size', type=int, default=4500, help='高分辨率预览长边像素')
    high.set_defaults(func=detail)
    args = parser.parse_args()
    if args.command == 'prepare' and (args.batch_size <= 0 or args.preview_size < 600):
        parser.error('--batch-size 必须大于0，--preview-size 必须 >=600')
    if args.command == 'detail' and args.size < 600:
        parser.error('detail --size 必须 >=600')
    try:
        args.func(args)
    except (ValueError, OSError, KeyError, RuntimeError, json.JSONDecodeError) as e:
        print('错误：' + str(e), file=sys.stderr)
        sys.exit(2)


if __name__ == '__main__':
    main()
